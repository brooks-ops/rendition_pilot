"""Account-list preview and one-PDF-per-account packet batches."""
import base64
import binascii
import csv
from io import BytesIO,StringIO
from pathlib import Path
import re
import zipfile

MAX_ACCOUNTS=100
MAX_FILE_BYTES=2*1024*1024
ACCOUNT=re.compile(r'R\s*([0-9]+)',re.I)
HEADERS={'account','accounts','accountnumber','accountnumbers','accountno','accountid','raccount','rnumber','cadaccount','propertyid','propid','propertyaccount','cadaccountnumber','accountnum','raccountnumber'}


def extract_cells(payload):
    filename=str(payload.get('filename') or '')
    if payload.get('file_base64'):
        encoded=payload['file_base64']
        if not isinstance(encoded,str) or len(encoded)>MAX_FILE_BYTES*4/3+8:raise ValueError('Account list files must be 2 MB or smaller.')
        try:data=base64.b64decode(encoded,validate=True)
        except (ValueError,binascii.Error):raise ValueError('The uploaded file could not be read.')
        if len(data)>MAX_FILE_BYTES:raise ValueError('Account list files must be 2 MB or smaller.')
        extension=Path(filename).suffix.lower()
        if extension=='.xlsx':
            from openpyxl import load_workbook
            try:
                with zipfile.ZipFile(BytesIO(data)) as archive:
                    if len(archive.infolist())>1000 or sum(i.file_size for i in archive.infolist())>20*1024*1024:raise ValueError('Excel file is too large after decompression.')
                workbook=load_workbook(BytesIO(data),read_only=True,data_only=True,keep_links=False)
                try:
                    if len(workbook.worksheets)>50:raise ValueError('Use an Excel file with 50 sheets or fewer.')
                    choices=[]
                    for sheet in workbook.worksheets:
                        sample=list(sheet.iter_rows(max_row=20,max_col=50,values_only=True))
                        col,header=choose_column(sample)
                        score=sum(bool(ACCOUNT.fullmatch(str(row[col] or '').strip())) for row in sample if len(row)>col)
                        choices.append((header is not None,score,sheet,col,header))
                    _,_,sheet,col,header=max(choices,key=lambda x:(x[0],x[1]))
                    cells=[]
                    for rownum,row in enumerate(sheet.iter_rows(min_col=col+1,max_col=col+1,values_only=True),1):
                        if rownum>10000:raise ValueError('Account lists must contain 10,000 rows or fewer.')
                        if header is not None and rownum<=header+1:continue
                        if row[0] is not None and str(row[0]).strip():cells.append((rownum,str(row[0]).strip()))
                    return cells,f'{filename} / {sheet.title} / column {col+1}'
                finally:workbook.close()
            except ValueError:raise
            except Exception as exc:raise ValueError('Unable to read this Excel file. Upload a valid .xlsx account list.') from exc
        if extension not in ('.csv','.txt'):raise ValueError('Upload a CSV, TXT or XLSX account list.')
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:text=data.decode('cp1252',errors='replace')
    else:
        text=payload.get('text','')
        if not isinstance(text,str) or len(text)>MAX_FILE_BYTES:raise ValueError('Account list is too large.')
        extension='.txt'
    if extension=='.csv':
        try:
            dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
        except csv.Error:dialect=csv.excel
        rows=list(csv.reader(StringIO(text),dialect))
        if len(rows)>10000:raise ValueError('Account lists must contain 10,000 rows or fewer.')
        if all(ACCOUNT.fullmatch(cell.strip()) for row in rows for cell in row if cell.strip()):
            return [(n+1,cell.strip()) for n,row in enumerate(rows) for cell in row if cell.strip()],filename
        col,header=choose_column(rows[:20])
        return [(n+1,row[col].strip()) for n,row in enumerate(rows) if len(row)>col and row[col].strip() and (header is None or n>header)],f'{filename} / column {col+1}'
    cells=[]
    for n,line in enumerate(text.splitlines(),1):
        # A plain list can be newline, comma, semicolon, tab or space separated.
        if re.sub(r'[^a-z]','',line.lower()) in HEADERS:continue
        line=re.sub(r'(?i)\bR[ \t]+(?=[0-9])','R',line)
        for value in re.split(r'[,;\t ]+',line.strip()):
            if value and re.sub(r'[^a-z]','',value.lower()) not in HEADERS:cells.append((n,value))
        if len(cells)>10000:raise ValueError('Account lists must contain 10,000 entries or fewer.')
    return cells,filename or 'Pasted account list'


def choose_column(rows):
    for n,row in enumerate(rows):
        for col,value in enumerate(row):
            if re.sub(r'[^a-z]','',str(value or '').lower()) in HEADERS:return col,n
    width=max(1,max((len(row) for row in rows),default=1))
    scores=[sum(bool(ACCOUNT.fullmatch(str(row[col] or '').strip())) for row in rows if len(row)>col) for col in range(width)]
    return max(range(width),key=lambda col:scores[col]),None


def preview_account_list(store,payload):
    cells,source=extract_cells(payload)
    if not cells:raise ValueError('No account entries found. Upload a list or paste R account numbers.')
    accounts=[];seen=set();duplicates=[];invalid=[]
    for line,value in cells:
        match=ACCOUNT.fullmatch(value)
        if not match:invalid.append({'row':line,'value':value[:200]});continue
        account='R'+match.group(1)
        if account in seen:duplicates.append(account);continue
        seen.add(account);accounts.append(account)
    if len(accounts)>MAX_ACCOUNTS:raise ValueError(f'Import up to {MAX_ACCOUNTS} distinct accounts per batch. Split this list into smaller batches.')
    matches={account:[] for account in accounts}
    _,rows,_=store.read_view()
    for row in rows:
        if row.get('reference_only'):continue
        account=row['account'].strip().upper()
        if account in matches:matches[account].append(row)
    return {'source':source,'account_count':len(accounts),'matched_count':sum(bool(matches[a]) for a in accounts),'model_rows':sum(map(len,matches.values())),'items':[{'account':a,'rows':matches[a]} for a in accounts if matches[a]],'unmatched':[a for a in accounts if not matches[a]],'duplicate_count':len(duplicates),'duplicates':list(dict.fromkeys(duplicates)),'invalid_count':len(invalid),'invalid':invalid[:100]}


def individual_packet_zip(store,ids,notes=''):
    from .packets import build_packet
    if not isinstance(ids,list) or not ids or len(ids)>500:raise ValueError('Select 1–500 model rows for an individual packet batch.')
    if any(not isinstance(rid,str) for rid in ids):raise ValueError('Invalid queued property IDs.')
    if not isinstance(notes,str) or len(notes)>5000:raise ValueError('Packet notes must contain 5,000 characters or fewer.')
    with store.connect() as db:revision=db.execute('SELECT revision FROM state').fetchone()[0]
    lookup={r['id']:r for r in store.records()};groups={}
    for rid in dict.fromkeys(ids):
        raw=lookup.get(rid)
        if not raw or raw.get('reference_only'):raise ValueError('A queued property is unavailable. Refresh the queue before generating packets.')
        account=raw['account'].strip().upper()
        if not re.fullmatch(r'R[0-9]+',account):raise ValueError('A queued property has an invalid account number.')
        groups.setdefault(account,[]).append(rid)
    if len(groups)>MAX_ACCOUNTS:raise ValueError('Generate up to 100 individual account packets per batch.')
    if any(len(rows)>100 for rows in groups.values()):raise ValueError('An account has more than 100 model rows. Select fewer sections for that account.')
    stream=BytesIO();manifest=StringIO();writer=csv.writer(manifest);writer.writerow(['Account','PDF file','Model sections','Source rows'])
    with zipfile.ZipFile(stream,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        for account,rows in groups.items():
            filename=f'{account}.pdf'
            try:pdf=build_packet(store,rows,notes=notes)
            except Exception as exc:raise ValueError(f'Unable to generate the packet for {account}. No partial ZIP was returned; review this property and try again.') from exc
            archive.writestr(filename,pdf)
            writer.writerow([account,filename,'; '.join(lookup[r]['model'] for r in rows),'; '.join(str(lookup[r]['source_row']) for r in rows)])
        archive.writestr('packet-manifest.csv',manifest.getvalue())
    with store.connect() as db:current_revision=db.execute('SELECT revision FROM state').fetchone()[0]
    if current_revision!=revision:raise ValueError('The workspace changed while packets were being generated. Generate the batch again to use consistent saved values.')
    return stream.getvalue()
