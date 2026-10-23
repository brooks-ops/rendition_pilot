import json
import gzip
from pathlib import Path
import pytest
from app.commercial.formulas import Engine,FormulaError,parse
from app.commercial.store import Store,Conflict


def test_formula_recalculation_uses_inputs_and_lazy_if():
    engine=Engine({'Model':{'A1':{'v':100},'B1':{'v':.1},'C1':{'v':1000,'f':'IF(B1=0,1/0,A1/B1)'},'D1':{'v':1010,'f':'ROUND(C1+10,0)'}}},{('Model','B1'):.2})
    assert engine.cell('Model','D1')==510
    assert Engine({'S':{'A1':{'v':-2},'B1':{'v':4,'f':'-A1^2'}}},{('S','A1'):-3}).cell('S','B1')==9


def test_lookup_ignores_errors_outside_the_matched_return_cell():
    cells={'A1':{'v':'R1'},'B1':{'v':100},'C1':{'v':'#REF!','error':True},'A2':{'v':'R2'},'B2':{'v':200},'D1':{'v':100,'f':'VLOOKUP(A1,A1:C2,2,FALSE)'},'E1':{'v':5,'f':'IFERROR(B1/0,5)'}}
    e=Engine({'S':cells},{('S','B1'):150})
    assert e.cell('S','D1')==150
    assert e.cell('S','E1')==5


def test_iferror_does_not_hide_unvalidated_calculations():
    e=Engine({'S':{'A1':{'v':1},'B1':{'v':0,'f':'IFERROR(UNSUPPORTED(A1),0)'}}},{('S','A1'):2})
    with pytest.raises(FormulaError,match='needs validation'):e.cell('S','B1')


def test_unavailable_reference_does_not_turn_into_zero():
    e=Engine({'S':{'A1':{'v':2},'B1':{'v':9,'f':"A1+'[1]Missing'!A1"}}},{('S','A1'):3})
    with pytest.raises(FormulaError):e.cell('S','B1')


@pytest.fixture
def store(tmp_path):
    sheet={'A1':{'v':'Account'},'B1':{'v':'Category'},'C1':{'v':'Cap Rate'},'D1':{'v':'Net operating income'},'E1':{'v':'Value'},'F1':{'v':'Area'}}
    records=[]
    for row,account in [(2,'R1'),(3,'R2')]:
        sheet.update({f'A{row}':{'v':account},f'B{row}':{'v':'Office A'},f'C{row}':{'v':.1},f'D{row}':{'v':100},f'E{row}':{'v':1000,'f':f'D{row}/C{row}'},f'F{row}':{'v':200}})
        records.append({'id':account,'model':'office','account':account,'source_sheet':'Office','source_row':row,'source_file':'office.xlsx','source_hash':'abc'})
    model={'model':'office','sheet':'Office','fields':{'type':'office','account':'A','category':'B','cap_rate':'C','noi':'D','total_value':'E','area_sq_ft':'F'}}
    (tmp_path/'manifest.json').write_text(json.dumps({'models':[model],'records':records,'graphs':{'office':'office.json.gz'}}))
    with gzip.open(tmp_path/'office.json.gz','wt') as f:json.dump({'Office':sheet},f)
    return Store(tmp_path)


def preview(store,**kwargs):
    return store.preview({'filters':{'model':'office','category':'Office A'},'field':'cap_rate','value':.2,**kwargs})


def test_preview_apply_undo_and_stale_token(store):
    p=preview(store);assert p['count']==2
    assert p['items'][0]['after_value']==500
    assert store.detail('R1')['values']['total_value']==1000
    batch=store.commit(p['token'],'appraiser','Review cap rates')
    assert store.detail('R1')['values']['total_value']==500
    with pytest.raises(Conflict):store.commit(p['token'],'appraiser','Repeated click')
    store.undo(batch['batch_id'],'appraiser')
    assert store.detail('R1')['values']['total_value']==1000
    assert store.catalog({'ids':['R1']})['items'][0]['values']['total_value']==1000
    assert store.catalog({'ids':['R1']})['items'][0]['values']['cap_rate']==.1


def test_keep_overrides_and_zero_values(store):
    p=preview(store,filters={'ids':['R1']});store.commit(p['token'],'appraiser','Property-specific rate')
    p=preview(store,value=.25)
    assert p['count']==1 and p['skipped_overrides']==1
    p=preview(store,field='area_sq_ft',value=0)
    store.commit(p['token'],'appraiser','Zero area is deliberate')
    assert store.detail('R2')['values']['area_sq_ft']==0


def test_revision_invalidates_preview_after_new_property(store):
    p=preview(store)
    created=store.create('R1',{'account':'R3','category':'Office A'},'appraiser')
    assert created['values']['account']=='R3'
    assert created['values']['total_value']==1000
    with pytest.raises(Conflict):store.commit(p['token'],'appraiser','Old preview')
    p=preview(store,value=.25,include_overrides=True)
    store.commit(p['token'],'appraiser','Category update includes new property')
    assert store.detail(created['id'])['values']['total_value']==400


def test_reject_invalid_cap_and_duplicate_model_account(store):
    with pytest.raises(ValueError):preview(store,value=0)
    with pytest.raises(Conflict):store.create('R1',{'account':'R2'},'appraiser')


def test_empty_selection_never_means_entire_model(store):
    with pytest.raises(ValueError):preview(store,filters={'ids':[]})


def test_preview_does_not_persist_authentication_token(store):
    preview(store,access_token='do-not-store-this-token')
    with store.connect() as db:
        assert 'do-not-store-this-token' not in db.execute('SELECT payload FROM previews').fetchone()[0]


def test_create_is_audited_and_can_be_undone(store):
    created=store.create('R1',{'account':'R3'},'appraiser')
    history=store.history()
    assert history[0]['reason']=='Add R3 to Office'
    store.undo(history[0]['id'],'appraiser')
    with pytest.raises(KeyError):store.detail(created['id'])


def test_new_property_inherits_effective_template_inputs(store):
    p=preview(store,filters={'ids':['R1']});store.commit(p['token'],'appraiser','Specific cap')
    created=store.create('R1',{'account':'R3'},'appraiser')
    assert created['values']['cap_rate']==.2
    assert created['values']['total_value']==500


def test_api_requires_district_context_for_every_data_request(store,monkeypatch):
    from types import SimpleNamespace
    from fastapi import FastAPI,HTTPException
    from fastapi.testclient import TestClient
    from app.commercial.api import commercial_router
    monkeypatch.setenv('COMMERCIAL_DATA_DIR',str(store.directory.parent))
    calls=[]
    def resolve(token):
        calls.append(token)
        if token!='valid':raise HTTPException(401,'Sign in required')
        return SimpleNamespace(district_id=store.directory.name,email='appraiser')
    app=FastAPI();app.include_router(commercial_router(resolve));client=TestClient(app)
    assert client.post('/api/commercial/catalog',json={}).status_code==401
    response=client.post('/api/commercial/catalog',json={'access_token':'valid','district_id':'another-district'})
    assert response.status_code==200
    assert response.json()['accounts']==2
    assert client.post('/api/commercial/property',json={'access_token':'valid','id':'missing'}).status_code==404
    assert client.post('/api/commercial/equity',json={'id':'R1'}).status_code==401
    assert client.post('/api/commercial/equity',json={'access_token':'valid','id':'R1'}).json()['comparables'][0]['account']=='R2'
    assert calls==['','valid','valid','','valid']


def test_automatic_equity_excludes_subject_and_updates_with_saved_values(store):
    from app.commercial.equity import generate_equity
    e=generate_equity(store,'R1')
    assert [r['account'] for r in e['comparables']]==['R2']
    assert e['median_unit_value']==5
    assert e['status']=='ready'
    p=preview(store,filters={'ids':['R2']});store.commit(p['token'],'appraiser','Revised rate')
    assert generate_equity(store,'R1')['median_unit_value']==2.5


def test_automatic_equity_rejects_invalid_size_and_other_categories(store):
    from app.commercial.equity import generate_equity
    store.create('R1',{'account':'R3','category':'Different use'},'appraiser')
    p=preview(store,filters={'ids':['R2']},field='area_sq_ft',value=0)
    store.commit(p['token'],'appraiser','Zero area')
    assert [r['account'] for r in generate_equity(store,'R1')['comparables']]==['R3']
    assert generate_equity(store,'R1')['comparables'][0]['match_tier']==3
    assert generate_equity(store,'R2')['subject']['unit_value'] is None


def test_automatic_equity_deduplicates_accounts_and_ranks_without_values(store):
    from app.commercial.equity import generate_equity
    store.create('R1',{'account':'R3'},'appraiser')
    p=preview(store,filters={'ids':['R2']},field='cap_rate',value=.01)
    store.commit(p['token'],'appraiser','Higher value does not affect ranking')
    e=generate_equity(store,'R1')
    assert [r['account'] for r in e['comparables']]==['R2','R3']
    assert e['comparables'][0]['comparison_value']==10000


def test_packets_always_include_automatic_equity(store):
    from app.commercial.packets import build_packet
    pdf=build_packet(store,['R1'])
    assert pdf.startswith(b'%PDF')


@pytest.mark.parametrize('model,basis',[('apartments','units'),('hotels','rooms'),('office','area_sq_ft')])
def test_equity_uses_model_basis_and_one_row_per_account(model,basis):
    from app.commercial.equity import generate_equity
    def row(id,account,size,category='A',classification='A'):
        return {'id':id,'account':account,'model':model,'source_row':2,'values':{'category':category,'classification':classification,basis:size,'total_value':1000},'errors':{}}
    records=[row('s','R1',100),row('subject-component','R1',100),row('a','R2',100),row('duplicate','R2',110),row('wrong-size','R3',301),row('wrong-class','R4',100,classification='B')]
    class FakeStore:
        def detail(self,id):return next(r for r in records if r['id']==id)
        def records(self):return records
        def snapshot(self,r):return r
    from app.commercial.equity import select_equity
    e=select_equity(records[0],records)
    assert e['basis']==basis
    assert [r['id'] for r in e['comparables']]==['a','wrong-size','wrong-class']
    assert e['comparables'][1]['match_tier']==1
    assert e['median_unit_value']==10


def test_equity_details_preserve_rates_zero_missing_and_errors():
    from app.commercial.equity import comparison_sections
    records=[{'values':{'cap_rate':.0625,'vacancy_rate':0,'expense_rate':.35,'noi':1234.56,'primary_rent':21,'projected_gross_rent':None},'errors':{'projected_gross_rent':'#REF!'}}, {'values':{'cap_rate':.07,'vacancy_rate':None,'expense_rate':.4,'noi':1000},'errors':{}}]
    sections=comparison_sections(records)
    rows={r['key']:r for s in sections for r in s['rows']}
    assert rows['cap_rate']['format']=='rate'
    assert rows['expense_rate']['cells'][1]['value']==.4
    assert rows['vacancy_rate']['cells'][0]['value']==0
    assert rows['vacancy_rate']['cells'][1]['value'] is None
    assert rows['noi']['format']=='money'
    assert rows['primary_rent']['label'].endswith('(source basis)')
    assert not rows['projected_gross_rent']['label'].endswith('(source basis)')
    assert rows['projected_gross_rent']['cells'][0]['error']=='#REF!'


def test_equity_always_fills_peers_for_missing_subject_and_sparse_category():
    from app.commercial.equity import select_equity
    subject={'id':'s','account':'R1','model':'office','source_row':2,'values':{'category':'Only property'},'errors':{}}
    pool=[subject]+[{'id':str(i),'account':'R'+str(i),'model':'office','source_row':i,'values':{'category':'Office B','leasable_area_sq_ft':100*i,'total_value':10000*i},'errors':{}} for i in range(2,10)]
    e=select_equity(subject,pool)
    assert len(e['comparables'])==6
    assert e['subject']['unit_value'] is None
    assert e['unadjusted_indication'] is None
    assert all(r['match_tier']==3 for r in e['comparables'])
    assert e['status']=='review'


def test_remote_storage_persists_changes_and_rolls_back_failed_writes(store,monkeypatch):
    import requests
    from fastapi import HTTPException
    from app.commercial.remote import RemoteStore
    remote=RemoteStore(store.directory,'district-one');uploads=[]
    class OK:
        def raise_for_status(self):pass
    def upload(method,district,name,data=None):
        uploads.append((method,district,name,data));return OK()
    monkeypatch.setattr('app.commercial.remote.object_request',upload)
    p=remote.preview({'filters':{'ids':['R1']},'field':'cap_rate','value':.2})
    b=remote.commit(p['token'],'appraiser','Durable rate')
    assert remote.detail('R1')['values']['total_value']==500
    assert uploads[0][:3]==('POST','district-one','workspace.sqlite3')
    assert uploads[0][3].startswith(b'SQLite format 3')
    remote.undo(b['batch_id'],'appraiser')
    assert remote.detail('R1')['values']['total_value']==1000
    def failed(*args):raise requests.ConnectionError('offline')
    monkeypatch.setattr('app.commercial.remote.object_request',failed)
    p=remote.preview({'filters':{'ids':['R1']},'field':'cap_rate','value':.3})
    with pytest.raises(HTTPException):remote.commit(p['token'],'appraiser','Must roll back')
    assert remote.detail('R1')['values']['total_value']==1000
    assert len(remote.history())==1


def test_remote_restore_uses_only_authenticated_district(store,monkeypatch,tmp_path):
    from app.commercial.remote import load_store
    import sqlite3
    import json
    backup=tmp_path/'snapshot.sqlite3'
    with store.connect() as source,sqlite3.connect(backup) as dest:source.backup(dest)
    objects={p.name:p.read_bytes() for p in store.directory.iterdir() if p.suffix in ('.json','.gz')}
    objects['workspace.sqlite3']=backup.read_bytes();calls=[]
    class Response:
        status_code=200
        def __init__(self,data):self.content=data
        def raise_for_status(self):pass
        def json(self):return json.loads(self.content)
    def download(method,district,name,data=None):
        calls.append((method,district,name));return Response(objects[name])
    monkeypatch.setenv('SUPABASE_SERVICE_ROLE_KEY','test-only')
    monkeypatch.setattr('app.commercial.remote.object_request',download)
    restored=load_store(tmp_path/'restored','district-one')
    assert restored.detail('R1')['values']['total_value']==1000
    assert all(district=='district-one' for _,district,_ in calls)


def test_commercial_uses_existing_rendition_account_and_district(monkeypatch):
    from fastapi.testclient import TestClient
    from types import SimpleNamespace
    from backend import main
    user={'id':'existing-user','email':'bbarrett@lubbockcad.org'}
    district=SimpleNamespace(district_id='existing-district',district_slug='lubbock',district_name='Lubbock CAD',domain='lubbockcad.org',email=user['email'],user_id=user['id'],role='admin')
    calls=[]
    def sign_in(email,password):
        calls.append(email);return {'access_token':'test-existing-session','user':user}
    monkeypatch.setattr(main,'sign_in_with_supabase',sign_in)
    monkeypatch.setattr(main,'get_supabase_user',lambda token:user)
    monkeypatch.setattr(main,'build_district_context',lambda user,token:district)
    client=TestClient(main.app)
    page=client.get('/commercial')
    assert page.status_code==200 and 'commercialLoginForm' in page.text
    assert '/api/auth/' in page.text and 'commercialSignOut' in page.text
    response=client.post('/api/auth/login',json={'email':'bbarrett@lubbockcad.org','password':'test-placeholder'})
    assert response.status_code==200
    assert response.json()['district']['district_id']=='existing-district'
    assert response.json()['user']['id']=='existing-user'
    restored=client.post('/api/auth/restore',json={'access_token':'test-existing-session'})
    assert restored.json()['district']==response.json()['district']
    assert calls==['bbarrett@lubbockcad.org']


def test_account_list_preview_matches_and_reports_exceptions(store):
    from app.commercial.packet_import import preview_account_list
    result=preview_account_list(store,{'text':'Account Numbers\nr1, R2\nR1\nR999\nnot-an-account'})
    assert [i['account'] for i in result['items']]==['R1','R2']
    assert result['unmatched']==['R999']
    assert result['duplicate_count']==1
    assert result['invalid']==[{'row':5,'value':'not-an-account'}]
    assert result['model_rows']==2


def test_account_list_reads_csv_column_and_headerless_lists(store):
    import base64
    from app.commercial.packet_import import preview_account_list
    csv='Owner,Account,Other\nExample,R2,R999\nSecond,R1,R888\n'
    result=preview_account_list(store,{'filename':'accounts.csv','file_base64':base64.b64encode(csv.encode()).decode()})
    assert [i['account'] for i in result['items']]==['R2','R1']
    assert result['unmatched']==[]
    result=preview_account_list(store,{'filename':'accounts.csv','file_base64':base64.b64encode(b'R1,R2').decode()})
    assert result['matched_count']==2


def test_account_list_reads_excel_sheet_and_account_header(store):
    import base64
    from io import BytesIO
    from openpyxl import Workbook
    from app.commercial.packet_import import preview_account_list
    workbook=Workbook();workbook.active.title='Instructions';workbook.active.append(['Use the accounts tab'])
    sheet=workbook.create_sheet('Accounts');sheet.append(['Owner','Account #']);sheet.append(['Example','R2']);sheet.append(['Example','R1']);sheet.append(['Other','R999'])
    stream=BytesIO();workbook.save(stream)
    result=preview_account_list(store,{'filename':'list.xlsx','file_base64':base64.b64encode(stream.getvalue()).decode()})
    assert result['matched_count']==2 and result['unmatched']==['R999']
    assert 'Accounts / column 2' in result['source']


def test_individual_zip_groups_all_selected_rows_per_account(store,monkeypatch):
    from app.commercial.packet_import import individual_packet_zip
    from io import BytesIO
    import zipfile
    raw=store.manifest['records'][0].copy();raw['id']='R1-component'
    with store.connect() as db:db.execute('INSERT INTO records(id,metadata) VALUES(?,?)',(raw['id'],json.dumps(raw)))
    calls=[]
    def build(s,ids,notes=''):
        calls.append(ids);return b'%PDF-test'
    monkeypatch.setattr('app.commercial.packets.build_packet',build)
    data=individual_packet_zip(store,['R1','R2','R1-component','R1'])
    with zipfile.ZipFile(BytesIO(data)) as archive:
        assert set(archive.namelist())=={'R1-ARB-packet.pdf','R2-ARB-packet.pdf','packet-manifest.csv'}
        assert archive.testzip() is None
        assert b'R1-ARB-packet.pdf' in archive.read('packet-manifest.csv')
    assert calls==[['R1','R1-component'],['R2']]


def test_individual_zip_has_real_separate_subject_pdfs(store):
    from app.commercial.packet_import import individual_packet_zip
    from io import BytesIO
    import zipfile,fitz
    with zipfile.ZipFile(BytesIO(individual_packet_zip(store,['R1','R2']))) as archive:
        for account,other in [('R1','R2'),('R2','R1')]:
            pdf=fitz.open(stream=archive.read(account+'-ARB-packet.pdf'),filetype='pdf')
            assert account in pdf[0].get_text() and other not in pdf[0].get_text()
            assert any('Equity comparison' in p.get_text() for p in pdf)


def test_packet_import_limits_and_invalid_batches(store):
    from app.commercial.packet_import import individual_packet_zip,preview_account_list
    with pytest.raises(ValueError,match='100 distinct'):preview_account_list(store,{'text':'\n'.join('R'+str(i) for i in range(101))})
    with pytest.raises(ValueError,match='No account'):preview_account_list(store,{'text':''})
    with pytest.raises(ValueError,match='could not be read'):preview_account_list(store,{'filename':'a.csv','file_base64':'not-base64'})
    with pytest.raises(ValueError,match='Select'):individual_packet_zip(store,[])
    with pytest.raises(ValueError,match='unavailable'):individual_packet_zip(store,['missing'])


def test_packet_zip_does_not_return_partial_or_mixed_revision_archive(store,monkeypatch):
    from app.commercial.packet_import import individual_packet_zip
    def fail(*args,**kwargs):raise RuntimeError('generation failure')
    monkeypatch.setattr('app.commercial.packets.build_packet',fail)
    with pytest.raises(ValueError,match='No partial ZIP'):individual_packet_zip(store,['R1','R2'])
    def changed(*args,**kwargs):
        with store.connect() as db:db.execute('UPDATE state SET revision=revision+1')
        return b'%PDF-test'
    monkeypatch.setattr('app.commercial.packets.build_packet',changed)
    with pytest.raises(ValueError,match='workspace changed'):individual_packet_zip(store,['R1'])


def test_account_packet_endpoints_auth_and_zip_delivery(store,monkeypatch):
    from types import SimpleNamespace
    from fastapi import FastAPI,HTTPException
    from fastapi.testclient import TestClient
    from app.commercial.api import commercial_router
    import zipfile
    from io import BytesIO
    monkeypatch.setenv('COMMERCIAL_DATA_DIR',str(store.directory.parent))
    def resolve(token):
        if token!='valid':raise HTTPException(401,'Sign in required')
        return SimpleNamespace(district_id=store.directory.name,email='appraiser')
    app=FastAPI();app.include_router(commercial_router(resolve));client=TestClient(app)
    for route in ('packet-list','packets/individual'):
        assert client.post('/api/commercial/'+route,json={}).status_code==401
    result=client.post('/api/commercial/packet-list',json={'access_token':'valid','text':'R1,R2,R999','district_id':'untrusted'}).json()
    assert result['matched_count']==2 and result['unmatched']==['R999']
    response=client.post('/api/commercial/packets/individual',json={'access_token':'valid','ids':['R1','R2']})
    assert response.status_code==200 and response.headers['content-type']=='application/zip'
    with zipfile.ZipFile(BytesIO(response.content)) as archive:assert len([n for n in archive.namelist() if n.endswith('.pdf')])==2
    assert client.post('/api/commercial/packets/individual',json={'access_token':'valid','ids':['missing']}).status_code==400


def test_multi_input_category_apply_and_undo(store):
    p=store.preview({'filters':{'model':'office'},'updates':[{'field':'cap_rate','value':.2},{'field':'area_sq_ft','operation':'percent','value':10}]})
    assert p['count']==2
    assert len(p['items'][0]['field_changes'])==2
    assert p['items'][0]['after_value']==500
    batch=store.commit(p['token'],'appraiser','Category review')
    assert batch['updated']==2
    assert store.detail('R1')['values']['area_sq_ft']==pytest.approx(220)
    assert store.detail('R1')['values']['cap_rate']==.2
    store.undo(batch['batch_id'],'appraiser')
    assert store.detail('R1')['values']['area_sq_ft']==200
    assert store.detail('R1')['values']['cap_rate']==.1


def test_category_summary_tracks_scope_and_overrides(store):
    from app.commercial.category_updates import category_summary
    p=preview(store,filters={'ids':['R1']});store.commit(p['token'],'appraiser','Test')
    d=category_summary(store,{'model':'office','ids':['R1']})
    assert d['count']==d['accounts']==1
    cap=next(f for f in d['fields'] if f['key']=='cap_rate')
    assert cap['median']==.2 and cap['overrides']==1 and cap['unit']=='%'
    assert not any(f['key']=='total_value' for f in d['fields'])
    p=store.preview({'filters':{'model':'office'},'updates':[{'field':'cap_rate','value':.3},{'field':'area_sq_ft','value':300}]})
    assert p['skipped_overrides']==1
    store.commit(p['token'],'appraiser','Protected multi edit')
    assert store.detail('R1')['values']['cap_rate']==.2
    assert store.detail('R1')['values']['area_sq_ft']==300


def test_multi_input_invalid_change_is_atomic(store):
    with pytest.raises(ValueError):store.preview({'updates':[{'field':'area_sq_ft','value':300},{'field':'cap_rate','value':-1}]})
    assert store.detail('R1')['values']['area_sq_ft']==200
    assert store.history()==[]
